import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as OpenpyxlImage
import os

wb = openpyxl.Workbook()

# Define styles
font_title = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="475569")
font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
font_row = Font(name="Segoe UI", size=9.5)
font_bold = Font(name="Segoe UI", size=9.5, bold=True)

fill_header_main = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep Navy
fill_header_sub = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid") # Ocean Blue
fill_header_gold = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid") # Amber Gold
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Slate 50
fill_resolved = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light Green

font_resolved = Font(name="Segoe UI", size=9.5, bold=True, color="166534") # Dark Green

thin_border_side = Side(border_style="thin", color="CBD5E1")
border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

user_uploaded_dir = r"C:\Users\Legion\.gemini\antigravity-ide\brain\d3efb056-1a87-4652-9f23-2a17db4ebdab\.user_uploaded"

# ==========================================
# SHEET 1: CHECKLIST NGHIỆP VỤ DÀNH RIÊNG CHO CHỊ DINH TINH
# ==========================================
ws1 = wb.active
ws1.title = "Dành Cho Khách Check"
ws1.views.sheetView[0].showGridLines = True

ws1.merge_cells("A1:G1")
ws1["A1"] = "BẢNG KIỂM TRA TÍNH NĂNG & BIỂU MẪU ĐÃ CHỈNH SỬA (USER ACCEPTANCE CHECKLIST)"
ws1["A1"].font = font_title
ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws1.row_dimensions[1].height = 28

ws1.merge_cells("A2:G2")
ws1["A2"] = "Dành riêng cho Chị Dinh Tinh kiểm tra thực tế trên phần mềm. Không dùng thuật ngữ kỹ thuật / code."
ws1["A2"].font = font_subtitle
ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws1.row_dimensions[2].height = 20

headers_ws1 = [
    "STT", 
    "Hình Ảnh Yêu Cầu Gốc",
    "Yêu Cầu Chị Đã Góp Ý", 
    "Vị Trí Cần Vào Kiểm Tra", 
    "Cách Thao Tác Thử Nghiệm", 
    "Kết Quả Thực Tế Đã Đạt Được",
    "Chị Xác Nhận (Đạt / Chưa)"
]

ws1.append([]) # Row 3 blank
ws1.append(headers_ws1) # Row 4
ws1.row_dimensions[4].height = 30

for col_num in range(1, len(headers_ws1) + 1):
    c = ws1.cell(row=4, column=col_num)
    c.font = font_header
    c.fill = fill_header_main
    c.alignment = align_center
    c.border = border_cell

checklist_items = [
    (
        1,
        os.path.join(user_uploaded_dir, "media_1788533937672.png"),
        "Xuất file Word (.docx) Hợp đồng, Phiếu thu, BEO bị lỗi, không xuất được để kiểm tra.",
        "Màn hình Hợp Đồng, Biên Nhận Cọc, BEO, Phụ Lục",
        "Chọn bất kỳ 1 Hợp đồng, Phiếu thu hoặc BEO -> Bấm nút 'Xuất DOCX' (In biểu mẫu).",
        "Tải file Word về máy mở lên mượt mà 100%, không còn bị báo lỗi file hay văng trang trắng.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        2,
        os.path.join(user_uploaded_dir, "media_1788533937672.png"),
        "Chia 2 mẫu Quyết toán: QT01 (cho tiệc Cưới, Cá nhân) và QT02 (cho tiệc Công ty, Pháp nhân).",
        "Nút 'Quyết Toán' trên danh sách Hợp Đồng",
        "Thử bấm Quyết toán cho 1 tiệc Cưới -> Xuất Word. Sau đó thử bấm Quyết toán cho 1 tiệc Hội nghị/Công ty -> Xuất Word.",
        "Phần mềm tự động tải đúng mẫu QT01 cho tiệc cưới/cá nhân, và tự động tải đúng mẫu QT02 cho tiệc công ty/hội nghị.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        3,
        os.path.join(user_uploaded_dir, "media_1788533943556.png"),
        "Set menu cho phép đổi món và tính bù giá chênh lệch theo từng món đổi.",
        "Form Hợp Đồng -> Tab 'Đổi Món Bù Giá'",
        "Tại Hợp đồng có Set menu, mở tab Đổi món -> Chọn 1 món gốc cần đổi và 1 món mới thay thế -> Nhìn cột tiền bù.",
        "Hệ thống tự động tính ra tiền chênh lệch = (Đơn giá món mới - Đơn giá món cũ) x Số lượng bàn tiệc.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        4,
        os.path.join(user_uploaded_dir, "media_1788533943556.png"),
        "Quyết toán tiệc phải lấy theo BEO sau lần thay đổi cuối cùng + Chi phí phát sinh tại tiệc.",
        "Màn hình Quyết Toán Tiệc",
        "Tạo 1 Hợp đồng có lập Phụ lục thay đổi BEO và có nhập phát sinh -> Bấm nút 'Quyết Toán'.",
        "Màn hình quyết toán tự động gom đúng số bàn/thực đơn của BEO mới nhất + tiền phát sinh + tiền đổi món.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        5,
        os.path.join(user_uploaded_dir, "media_1788534140892.png"),
        "Mẫu trang trí không dán hình cố định, tạo thư mục ảnh để tick chọn mẫu cho từng tiệc.",
        "Form Hợp Đồng -> Bấm nút '🎨 Mẫu Trang Trí'",
        "Khi tạo/sửa Hợp đồng, bấm nút 'Mẫu Trang Trí' ở góc trên -> Xem gallery các mẫu hoa, backdrop kèm giá -> Bấm 'Chọn Mẫu Này'.",
        "Có thư viện hình ảnh trực quan để xem trước và tick chọn; Tên mẫu và đơn giá trang trí tự động điền vào Hợp đồng/BEO.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        6,
        os.path.join(user_uploaded_dir, "media_1788534274024.png"),
        "HĐ Triển lãm + Tiệc: Địa điểm triển lãm là 1 sảnh, địa điểm tiệc là 1 sảnh riêng (tiệc trưa hoặc tối linh hoạt).",
        "Form Hợp Đồng -> Tab 'Sảnh Tiệc'",
        "Chọn loại tiệc Triển lãm + Tiệc: Chọn 1 sảnh làm Triển Lãm, chọn thêm 1 sảnh khác làm Tiệc (chọn ca Trưa hoặc Tối) -> Xuất Word.",
        "File Word in ra hiển thị riêng biệt: 'Địa điểm triển lãm: Sảnh A' và 'Địa điểm tiệc (Trưa/Tối): Sảnh B'. Không bị gán cứng tên sảnh.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        7,
        os.path.join(user_uploaded_dir, "media_1788534269431.png"),
        "Thuê sảnh phát sinh ngoài giờ: lấy theo đơn giá sảnh, công thức tính = (Giá thuê 1 buổi 4 tiếng) / 4.",
        "Hợp Đồng Triển Lãm & Hội Nghị -> Điều 2",
        "Mở Hợp đồng triển lãm/hội nghị và xuất file Word -> Xem dòng 'Thuê sảnh phát sinh ngoài thời gian hợp đồng'.",
        "Đã bỏ con số cứng 17.500.000 VNĐ. Giá ngoài giờ tự động chia 4 từ đơn giá sảnh (Ví dụ sảnh 40 triệu/buổi -> in ra 10 triệu/giờ/sảnh).",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        8,
        os.path.join(user_uploaded_dir, "media_1788534274024.png"),
        "Số hợp đồng không được mặc định sẵn trong mẫu in, để mở tự nhập theo thực tế.",
        "Màn hình Hợp Đồng & Bản in Word",
        "Nhập số hợp đồng thực tế (Ví dụ: HD-2026/001) -> Bấm Lưu và Xuất Word.",
        "Toàn bộ số hợp đồng mẫu bị in dính trước đây đã được xóa sạch; trên file Word hiển thị đúng số hợp đồng chị vừa nhập.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        9,
        os.path.join(user_uploaded_dir, "media_1788534278166.png"),
        "Hợp đồng Triển Lãm riêng: sửa câu chữ đúng cho thuê mặt bằng triển lãm, bỏ hoàn toàn chữ 'và tiệc'.",
        "Chọn loại tiệc 'Triển Lãm' -> Xuất Hợp Đồng (Mẫu 2.2)",
        "Xuất file Word Hợp đồng Triển lãm và đọc đoạn đầu (Căn cứ & Điều 1).",
        "Đã gọt sạch chữ 'và tiệc'. Văn bản in chuẩn: 'Hợp đồng tổ chức triển lãm... Bên B có nhu cầu thuê địa điểm để tổ chức triển lãm...'.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        10,
        os.path.join(user_uploaded_dir, "media_1788534282258.png"),
        "Kích thước sân khấu bị in lặp 2 lần diện tích liên tiếp trong file Word.",
        "Mục Kích thước sân khấu trong HĐ Triển Lãm & Hội Nghị",
        "Xuất file Word Hợp đồng Triển Lãm hoặc Hội Nghị -> Xem phần thông số kỹ thuật sảnh.",
        "Đã sửa triệt để: Kích thước sân khấu chỉ hiển thị đúng 1 lần (Ví dụ: 9.6m x 3.6m x 0.8m), không còn bị nhân đôi chữ.",
        " [  ] Đạt   [  ] Chưa"
    ),
    (
        11,
        os.path.join(user_uploaded_dir, "media_1788534282258.png"),
        "Khách pháp nhân: Ký HĐ xong mới cọc; Không mặc định 70% tiền cọc mà để mở tự nhập tùy thỏa thuận (30% - 40%).",
        "Mục Đặt cọc trong 4 mẫu HĐ Công ty / Hội Nghị / Triển Lãm",
        "Nhập tỷ lệ cọc đợt 1 là 30% hoặc 40% trên form HĐ -> Xuất file Word hợp đồng.",
        "Đã xóa vĩnh viễn chữ '70%' cố định. File Word in đúng câu: 'Trong thời hạn 03 ngày làm việc kể từ ngày ký HĐ... đặt cọc số tiền: [Số tiền thỏa thuận]'.",
        " [  ] Đạt   [  ] Chưa"
    )
]

for row_idx, item in enumerate(checklist_items, start=5):
    stt, img_path, fb, loc, act, res, sign = item
    ws1.append([stt, "", fb, loc, act, res, sign])
    ws1.row_dimensions[row_idx].height = 100
    
    c_stt = ws1.cell(row=row_idx, column=1)
    c_stt.alignment = align_center
    c_stt.font = font_bold
    c_stt.border = border_cell
    
    c_img = ws1.cell(row=row_idx, column=2)
    c_img.border = border_cell
    c_img.alignment = align_center
    
    if os.path.exists(img_path):
        try:
            img = OpenpyxlImage(img_path)
            aspect = img.width / img.height
            if aspect > 1.6:
                img.width = 150
                img.height = int(150 / aspect)
            else:
                img.height = 90
                img.width = int(90 * aspect)
            cell_ref = f"B{row_idx}"
            ws1.add_image(img, cell_ref)
        except Exception as e:
            c_img.value = "[Ảnh]"
            
    c_fb = ws1.cell(row=row_idx, column=3)
    c_fb.alignment = align_left
    c_fb.font = font_bold
    c_fb.border = border_cell
    
    c_loc = ws1.cell(row=row_idx, column=4)
    c_loc.alignment = align_left
    c_loc.font = font_bold
    c_loc.border = border_cell
    c_loc.fill = fill_zebra
    
    c_act = ws1.cell(row=row_idx, column=5)
    c_act.alignment = align_left
    c_act.font = font_row
    c_act.border = border_cell
    
    c_res = ws1.cell(row=row_idx, column=6)
    c_res.alignment = align_left
    c_res.font = font_row
    c_res.border = border_cell
    c_res.fill = fill_resolved
    
    c_sign = ws1.cell(row=row_idx, column=7)
    c_sign.alignment = align_center
    c_sign.font = font_bold
    c_sign.border = border_cell

col_widths_ws1 = [6, 24, 34, 28, 44, 44, 20]
for idx, width in enumerate(col_widths_ws1, start=1):
    ws1.column_dimensions[get_column_letter(idx)].width = width


# ==========================================
# SHEET 2: HƯỚNG DẪN BẰNG NGÔN NGỮ NGHIỆP VỤ (KHÔNG DÙNG TỪ KỸ THUẬT)
# ==========================================
ws2 = wb.create_sheet(title="Hướng Dẫn Chị Tinh Thử Nghiệm")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:E1")
ws2["A1"] = "HƯỚNG DẪN CÁC BƯỚC THỬ NGHIỆM PHẦN MỀM TỪ ĐẦU ĐẾN CUỐI"
ws2["A1"].font = font_title
ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[1].height = 28

ws2.merge_cells("A2:E2")
ws2["A2"] = "Các bước thao tác thực tế dành cho người dùng vận hành, kế toán và quản lý sảnh."
ws2["A2"].font = font_subtitle
ws2["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws2.row_dimensions[2].height = 20

headers_ws2 = [
    "Thứ Tự", 
    "Giai Đoạn Nghiệp Vụ", 
    "Vào Màn Hình Nào?", 
    "Chị Thử Làm Những Gì?", 
    "Phần Mềm Sẽ Hiển Thị Kết Quả Thế Nào?"
]

ws2.append([])
ws2.append(headers_ws2)
ws2.row_dimensions[4].height = 30

for col_num in range(1, len(headers_ws2) + 1):
    c = ws2.cell(row=4, column=col_num)
    c.font = font_header
    c.fill = fill_header_sub
    c.alignment = align_center
    c.border = border_cell

steps_user = [
    (
        "Bước 1",
        "Tiếp khách & Đặt cọc giữ chỗ",
        "Menu 'Khách Tham Quan' -> Lập 'Biên Nhận Cọc Chỗ'",
        "Nhập tên khách, ngày dự kiến tổ chức tiệc, chọn sảnh và số tiền khách đặt cọc trước -> Bấm Lưu -> Bấm 'In Phiếu Thu' (Xuất Word).",
        "File Word Phiếu Thu tải về máy, mở lên hiển thị đầy đủ tên khách, số tiền cọc bằng số & bằng chữ, không bị lỗi font hay trắng trang."
    ),
    (
        "Bước 2",
        "Lập Hợp Đồng Tiệc Cưới",
        "Menu 'Hợp Đồng' -> Bấm 'Thêm Mới' (Loại tiệc: Tiệc Cưới)",
        "1. Chọn thực đơn (Set menu hoặc chọn món lẻ).\n2. Bấm nút 'Mẫu Trang Trí' ở trên để xem ảnh và tick chọn mẫu hoa/backdrop yêu thích.\n3. Nếu ngày tiệc gấp (dưới 30 ngày), hệ thống tự kích hoạt mẫu chọn ngay menu.\n4. Bấm 'Lưu' và bấm 'Xuất DOCX'.",
        "File Word Hợp đồng in ra chuẩn đẹp: Tên công ty nhà hàng lấy chuẩn, phí phục vụ hiển thị linh hoạt, có thông tin mẫu trang trí khách vừa chọn."
    ),
    (
        "Bước 3",
        "Lập Hợp Đồng Công Ty / Triển Lãm / Hội Nghị",
        "Menu 'Hợp Đồng' -> Thêm mới (Loại: Triển Lãm / Hội Nghị)",
        "1. Chọn sảnh Triển Lãm riêng, chọn thêm 1 sảnh Tiệc riêng (chọn ca trưa hoặc tối).\n2. Nhập số tiền đặt cọc theo thỏa thuận (30% hoặc 40%).\n3. Bấm Lưu và bấm 'Xuất DOCX'.",
        "File Word mẫu 2.1 / 2.2 / 3.1 / 3.2 in ra đúng văn phong cho thuê địa điểm (bỏ chữ 'và tiệc' nếu chỉ thuê triển lãm), không bị lặp kích thước sân khấu, không dán cứng 70% cọc, giá ngoài giờ = giá sảnh chia 4."
    ),
    (
        "Bước 4",
        "Khách muốn đổi món trong Set menu",
        "Mở Hợp Đồng -> Bấm vào tab 'Đổi Món Bù Giá'",
        "Chọn 1 món gốc trong Set cần bỏ ra, chọn 1 món mới muốn đổi vào -> Nhìn số tiền phần mềm tự tính.",
        "Hệ thống tự tính: Tiền bù = (Giá món mới - Giá món gốc) x Số bàn. Khoản tiền này được lưu lại để lát nữa đưa vào quyết toán."
    ),
    (
        "Bước 5",
        "Thay đổi BEO trước ngày tiệc",
        "Menu 'Phụ Lục / Đề Nghị Thay Đổi BEO'",
        "Điều chỉnh tăng/giảm số bàn, đổi thêm món ăn hoặc dịch vụ -> Bấm 'Xuất BEO'.",
        "File BEO in ra hiển thị rõ ràng số bàn mặn, số bàn chay, thực đơn có ghi chú rõ món nào là món chay."
    ),
    (
        "Bước 6",
        "Ghi nhận phát sinh thực tế trong tiệc",
        "Tại dòng Hợp Đồng -> Mở tab 'Phát Sinh'",
        "Nhập các chi phí phát sinh lúc tiệc diễn ra (nước ngọt, bia gọi thêm, bàn phát sinh, giờ phụ thu...).",
        "Lưu lại đầy đủ các món/dịch vụ phát sinh thực tế vào hồ sơ tiệc."
    ),
    (
        "Bước 7",
        "Quyết toán tiệc",
        "Chọn Hợp Đồng trên danh sách -> Bấm nút 'Quyết Toán'",
        "Màn hình Quyết toán mở lên: Chị kiểm tra xem đã gom đủ số liệu BEO mới nhất + tiền bù đổi món + phát sinh thực tế - tiền cọc chưa -> Bấm 'Lưu & In'.",
        "Tự động tải về file Word QT01 (nếu là tiệc cưới/cá nhân) hoặc QT02 (nếu là công ty/hội nghị). Mọi con số cộng trừ hiển thị rõ ràng, chuẩn chỉnh!"
    )
]

for row_idx, item in enumerate(steps_user, start=5):
    stt, stage, scr, act, exp = item
    ws2.append([stt, stage, scr, act, exp])
    ws2.row_dimensions[row_idx].height = 80
    
    c_stt = ws2.cell(row=row_idx, column=1)
    c_stt.alignment = align_center
    c_stt.font = font_bold
    c_stt.border = border_cell
    c_stt.fill = fill_zebra
    
    c_stage = ws2.cell(row=row_idx, column=2)
    c_stage.alignment = align_center
    c_stage.font = font_bold
    c_stage.border = border_cell
    
    c_scr = ws2.cell(row=row_idx, column=3)
    c_scr.alignment = align_left
    c_scr.font = font_bold
    c_scr.border = border_cell
    
    c_act = ws2.cell(row=row_idx, column=4)
    c_act.alignment = align_left
    c_act.font = font_row
    c_act.border = border_cell
    
    c_exp = ws2.cell(row=row_idx, column=5)
    c_exp.alignment = align_left
    c_exp.font = font_row
    c_exp.border = border_cell

col_widths_ws2 = [10, 26, 30, 48, 48]
for idx, width in enumerate(col_widths_ws2, start=1):
    ws2.column_dimensions[get_column_letter(idx)].width = width


# ==========================================
# SHEET 3: CHI TIẾT KỸ THUẬT (DÀNH CHO DEV / IT NỘI BỘ THAM KHẢO)
# ==========================================
ws3 = wb.create_sheet(title="Kỹ Thuật Nội Bộ (IT Ref)")
ws3.views.sheetView[0].showGridLines = True

ws3.merge_cells("A1:E1")
ws3["A1"] = "BẢNG ĐỐI CHIẾU MÃ KỸ THUẬT NỘI BỘ (DÀNH CHO IT / DEV)"
ws3["A1"].font = font_title
ws3["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws3.row_dimensions[1].height = 28

ws3.merge_cells("A2:E2")
ws3["A2"] = "Lưu vết các file SQL, hàm JavaScript và template Word tương ứng với từng yêu cầu."
ws3["A2"].font = font_subtitle
ws3["A2"].alignment = Alignment(horizontal="left", vertical="center")
ws3.row_dimensions[2].height = 20

headers_ws3 = [
    "Mã REQ", 
    "Nhóm Nghiệp Vụ", 
    "Yêu Cầu BA", 
    "Trạng Thái", 
    "Tệp Code / SQL / Template Liên Quan"
]

ws3.append([])
ws3.append(headers_ws3)
ws3.row_dimensions[4].height = 28

for col_num in range(1, len(headers_ws3) + 1):
    c = ws3.cell(row=4, column=col_num)
    c.font = font_header
    c.fill = fill_header_gold
    c.alignment = align_center
    c.border = border_cell

it_refs = [
    ("REQ-01", "DOCX Engine", "Khắc phục lỗi render Word", "RESOLVED", "backend-app/server.js, DocumentExportPlugin.js"),
    ("REQ-02", "Settlement", "2 mẫu quyết toán QT01/QT02", "RESOLVED", "quyet_toan_01.docx, quyet_toan_02.docx, server.js:350"),
    ("REQ-03, 04", "Food & Menu", "Đổi món bù giá theo món", "RESOLVED", "DoiMonPlugin.js, Update_dmMauTrangTri_DoiMon.sql, Update_QuyetToan_AllInOne.sql"),
    ("REQ-05, 06", "Settlement", "Quyết toán = BEO cuối + phát sinh", "RESOLVED", "Update_QuyetToan_AllInOne.sql, tbmk_HopdongPhatSinh"),
    ("REQ-07", "Decoration", "Gallery thư viện mẫu trang trí", "RESOLVED", "dmMauTrangTri, MauTrangTriPlugin.js"),
    ("REQ-08, 09", "Hall Roles", "Tách vai trò sảnh & ca tiệc", "RESOLVED", "tbmk_Hopdongsanhtiec(LoaiDiaDiem, Thoigianid), v_DanhSachHopDong.sql"),
    ("REQ-10", "Contract No", "Số HĐ không hard-code", "RESOLVED", "v_DanhSachHopDong.sql, các file DOCX"),
    ("REQ-11, 12", "Overtime Fee", "Giá ngoài giờ = dmSanhtiec / 4", "RESOLVED", "v_DanhSachHopDong.sql (PhiThueSanhNgoaiGio)"),
    ("REQ-13, 14", "Exhibition", "Wording HĐ triển lãm chuẩn", "RESOLVED", "2.2 MAU HDONG (TRIỂN LÃM).docx, v_DanhSachHopDong.sql"),
    ("REQ-15", "Stage Size", "Xóa lặp kích thước sân khấu", "RESOLVED", "2.2 MAU HDONG (TRIỂN LÃM).docx"),
    ("REQ-16, 17", "Deposit", "Bỏ 70% cọc, ký HĐ xong mới cọc", "RESOLVED", "Các mẫu 2.1, 2.2, 3.1, 3.2 DOCX"),
    ("REQ-37, 38", "Vegetarian", "Tách món mặn & chay, 6 trường số bàn", "RESOLVED", "fn_DOCX_MenuDichVu.sql, Update_PhuLuc_AllInOne.sql")
]

for row_idx, item in enumerate(it_refs, start=5):
    ws3.append(list(item))
    ws3.row_dimensions[row_idx].height = 26
    
    for c_idx in range(1, 6):
        cell = ws3.cell(row=row_idx, column=c_idx)
        cell.font = font_row
        cell.border = border_cell
        if c_idx == 1 or c_idx == 4:
            cell.alignment = align_center
            if c_idx == 4:
                cell.fill = fill_resolved
                cell.font = font_resolved
        else:
            cell.alignment = align_left

col_widths_ws3 = [12, 18, 38, 16, 55]
for idx, width in enumerate(col_widths_ws3, start=1):
    ws3.column_dimensions[get_column_letter(idx)].width = width

wb.save("BANG_THEO_DOI_CHINH_SUA_CHANGELOG_BA.xlsx")
print("User-friendly checklist generated successfully!")
